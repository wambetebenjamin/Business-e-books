#!/usr/bin/env python3
"""
Verification pass over the generated deck.

1. Re-opens the saved .pptx (does not re-use in-memory state).
2. Estimates rendered text height per text frame using DejaVuSans metrics as a
   conservative proxy for Calibri, and flags boxes whose text will overflow.
3. Checks coverage of all 13 paper sections, the 5 disturbances, the case study,
   citation budget, speaker notes, and that the reference list matches the paper.
"""
import re
import sys
from pptx import Presentation
from pptx.util import Emu

from deck_metrics import para_height_in

DECK = "Gestalt Theory and Its Application to Counselling (11).pptx"
PAPER = "/tmp/gw/groupwork.txt"

prs = Presentation(DECK)
paper = open(PAPER, encoding="utf-8").read()
fails = []
warns = []

print("=" * 78)
print("1 · GEOMETRY / OVERFLOW CHECK")
print("=" * 78)
for idx, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        if not tf.text.strip():
            continue
        box_w = Emu(sh.width).inches - Emu(tf.margin_left).inches - Emu(tf.margin_right).inches
        box_h = Emu(sh.height).inches - Emu(tf.margin_top).inches - Emu(tf.margin_bottom).inches
        total_h = 0.0
        for p in tf.paragraphs:
            runs = p.runs
            if not runs:
                total_h += 0.14
                continue
            size = max((r.font.size.pt if r.font.size else 14) for r in runs)
            bold = any(bool(r.font.bold) for r in runs)
            txt = "".join(r.text for r in runs)
            ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
            sa = p.space_after.pt if p.space_after else 0.0
            sb = p.space_before.pt if p.space_before else 0.0
            total_h += para_height_in(txt, size, box_w, bold, ls, sb, sa)
        over = total_h - box_h
        if over > 0.10:
            fails.append(f"slide {idx:2d}  OVERFLOW +{over:.2f}in  ({total_h:.2f}in text in {box_h:.2f}in box)  :: {tf.text[:62]!r}")
        elif over > -0.06:
            warns.append(f"slide {idx:2d}  TIGHT   {over:+.2f}in  :: {tf.text[:62]!r}")

for f in fails:
    print("  FAIL", f)
for w in warns:
    print("  warn", w)
print(f"  -> {len(fails)} overflow, {len(warns)} tight")

print()
print("=" * 78)
print("2 · COVERAGE OF THE PAPER'S 13 SECTIONS")
print("=" * 78)
alltext = "\n".join(
    sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
)
checks = {
    "1  Introduction": ["integrated whole", "experiential"],
    "2  History / contributors": ["Fritz Perls", "Laura Perls", "Paul Goodman", "1951"],
    "3.1 Phenomenology": ["Phenomenology"],
    "3.2 Existential philosophy": ["Existential"],
    "3.3 Holism": ["Holism"],
    "3.4 Field theory": ["Field theory"],
    "4  Core assumptions": ["integrated wholes", "organismic self-regulation"],
    "5.1 Awareness / here-and-now": ["here-and-now", "noticing"],
    "5.2 Figure and ground": ["Figure and Ground", "ground"],
    "5.3 Contact / boundaries": ["Contact Boundaries", "negotiated"],
    "5.4 Unfinished business": ["Unfinished business", "resentment"],
    "5.5 Organismic self-regulation": ["self-regulation", "autonomy"],
    "5.6 Cycle of experience": ["Cycle of Experience", "Sensation", "Mobilisation",
                               "Satisfaction", "Withdrawal"],
    "5.7 Personal responsibility": ["Personal Responsibility", "BLAME"],
    "6.1 Introjection": ["Introjection", "unquestioned internal rule"],
    "6.2 Projection": ["Projection", "attributing"],
    "6.3 Retroflection": ["Retroflection", "inward"],
    "6.4 Deflection": ["Deflection", "distraction"],
    "6.5 Confluence": ["Confluence", "distinguishing self from others"],
    "6.6 Comparison table": ["I must never disappoint anyone", "Whatever you want is what I want"],
    "6.7 Clinical significance": ["retaliation", "protective function"],
    "7  Goals / counsellor role": ["Therapeutic goals", "I am fine"],
    "8.1 Awareness exercises": ["Awareness exercises"],
    "8.2 Empty chair": ["Empty chair"],
    "8.3 Two chair": ["Two-chair"],
    "8.4 Staying with a feeling": ["Staying with a feeling"],
    "8.5 Body awareness / exaggeration": ["Body awareness", "exaggeration"],
    "8.6 Language / I-statements": ["first-person", "I feel worthless"],
    "9  Case study (Amina)": ["Amina", "administrative officer", "32"],
    "10 Strengths": ["Strengths of Gestalt Therapy", "Holistic understanding"],
    "11 Limitations": ["Limitations and Criticisms", "evidence base"],
    "12 Critical evaluation": ["Critical Evaluation"],
    "13 Conclusion": ["CONCLUSION", "flexibility and choice"],
}
for k, needles in checks.items():
    missing = [n for n in needles if n.lower() not in alltext.lower()]
    if missing:
        fails.append(f"COVERAGE {k}: missing {missing}")
        print(f"  FAIL {k:34s} missing {missing}")
    else:
        print(f"  ok   {k}")

print()
print("=" * 78)
print("3 · CITATION BUDGET")
print("=" * 78)
pat = re.compile(r"\((?:[A-Z][A-Za-z\-]+(?:\s*(?:&|et al\.)[^)]*)?,\s*\d{4}[a-z]?\s*(?:;\s*)?)+\)")
total_cites = 0
for idx, slide in enumerate(prs.slides, 1):
    txt = "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    found = pat.findall(txt)
    total_cites += len(found)
    flag = "  <-- MORE THAN ONE" if len(found) > 1 else ""
    if found:
        print(f"  slide {idx:2d}: {len(found)}  {found}{flag}")
    if len(found) > 1:
        fails.append(f"CITATIONS slide {idx} has {len(found)} inline citations")
print(f"  -> {total_cites} inline citation groups (was 17 in the old deck)")

print()
print("=" * 78)
print("4 · CITATION / REFERENCE INTEGRITY vs THE PAPER")
print("=" * 78)
paper_refs = {
    "Brownell, P. (Ed.). (2010)": "Brownell",
    "Joyce, P., & Sills, C. (2014)": "Joyce & Sills",
    "Perls, F. S., Hefferline, R. F., & Goodman, P. (1951)": "Perls et al.",
    "Wagner-Moore, L. E. (2004)": "Wagner-Moore",
    "Wheeler, G., & Axelsson, L. (2015)": "Wheeler & Axelsson",
    "Woldt, A. L., & Toman, S. M. (Eds.). (2005)": "Woldt & Toman",
}
PAPER_SURNAMES = {"Brownell", "Joyce", "Sills", "Perls", "Hefferline", "Goodman",
                  "Wagner-Moore", "Wheeler", "Axelsson", "Woldt", "Toman"}

ref_slide = prs.slides[-1]
ref_text = "\n".join(sh.text_frame.text for sh in ref_slide.shapes if sh.has_text_frame)
for ref in paper_refs:
    if ref in ref_text:
        print(f"  ok   reference present : {ref}")
    else:
        fails.append(f"REFERENCE missing from the references slide: {ref}")
        print(f"  FAIL reference missing : {ref}")

# every surname listed on the references slide must appear in the paper's list
deck_ref_surnames = set(re.findall(r"^([A-Z][A-Za-z\-]+),\s[A-Z]\.", ref_text, re.M))
stray_refs = deck_ref_surnames - PAPER_SURNAMES
if stray_refs:
    fails.append(f"REFERENCE LISTED ON SLIDE BUT NOT IN THE PAPER: {sorted(stray_refs)}")
    print(f"  FAIL not in paper      : {sorted(stray_refs)}")
else:
    print(f"  ok   references slide lists nothing outside the paper: {sorted(deck_ref_surnames)}")

# every author named in an in-slide citation must resolve to the paper's list
cited = set(re.findall(r"([A-Z][A-Za-z\-]+)(?:\s+et\s+al\.|\s*&\s*[A-Z][A-Za-z\-]+)?,\s*\d{4}", alltext))
stray = cited - PAPER_SURNAMES
if stray:
    fails.append(f"CITED AUTHOR NOT IN THE PAPER'S REFERENCE LIST: {sorted(stray)}")
    print(f"  FAIL cited, not in paper: {sorted(stray)}")
else:
    print(f"  ok   every in-slide citation resolves to the paper: {sorted(cited)}")
for bad in ("Corey", "Yontef"):
    if bad in alltext:
        fails.append(f"STALE CITATION still present: {bad}")
        print(f"  FAIL stale citation present: {bad}")
    else:
        print(f"  ok   stale citation removed: {bad}")

print()
print("=" * 78)
print("5 · SPEAKER NOTES")
print("=" * 78)
no_notes = []
short = []
for idx, slide in enumerate(prs.slides, 1):
    if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
        no_notes.append(idx)
    elif len(slide.notes_slide.notes_text_frame.text.split()) < 40:
        short.append(idx)
if no_notes:
    fails.append(f"SLIDES WITHOUT NOTES: {no_notes}")
print(f"  slides without notes : {no_notes or 'none'}")
print(f"  notes under 40 words : {short or 'none'}")
wc = [len(s.notes_slide.notes_text_frame.text.split()) for s in prs.slides]
print(f"  notes word count     : min {min(wc)}, max {max(wc)}, total {sum(wc)}")

print()
print("=" * 78)
print("6 · STRUCTURE")
print("=" * 78)
print(f"  slide count          : {len(prs.slides)}")
print(f"  slide size           : {Emu(prs.slide_width).inches} x {Emu(prs.slide_height).inches} in")
npics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
print(f"  raster pictures      : {npics}  (old deck: 43)")
import os
print(f"  file size            : {os.path.getsize(DECK)/1024:.0f} KB  (old deck: {os.path.getsize('Gestalt Theory and Its Application to Counselling (10).pptx')/1024:.0f} KB)")

# shapes must stay inside the slide canvas
W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
for idx, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        r = Emu(sh.left).inches + Emu(sh.width).inches
        b = Emu(sh.top).inches + Emu(sh.height).inches
        if Emu(sh.left).inches < -0.01 or Emu(sh.top).inches < -0.01 or r > W + 0.01 or b > H + 0.01:
            fails.append(f"slide {idx}: shape {sh.name!r} outside canvas "
                         f"(right {r:.2f} > {W:.2f}, bottom {b:.2f} > {H:.2f})")

# no text below the readability floor
small = []
for idx, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() and r.font.size and r.font.size.pt < 8.5:
                    small.append((idx, r.font.size.pt, r.text[:30]))
for s_ in small:
    fails.append(f"slide {s_[0]}: text below 8.5pt ({s_[1]}pt) :: {s_[2]!r}")
print(f"  shapes off-canvas    : {sum(1 for f in fails if 'outside canvas' in f)}")
print(f"  text below 8.5pt     : {len(small)}")

print()
print("=" * 78)
if fails:
    print(f"RESULT: {len(fails)} PROBLEM(S)")
    for f in fails:
        print("   -", f)
    sys.exit(1)
print("RESULT: ALL CHECKS PASSED")
